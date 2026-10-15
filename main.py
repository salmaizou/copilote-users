import csv
import warnings
from pathlib import Path


from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
from openpyxl.utils import get_column_letter


PROJECT_DIR = Path(__file__).resolve().parent
DATA_DIR = PROJECT_DIR / "data"
OUTPUT_FILE = DATA_DIR / "donnees_preparees.xlsx"
MISSING_VALUE = "non"
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384

SOURCES = (
    ("train.csv", "Entrainement", "TableEntrainement", False),
    ("test.csv", "Test", "TableTest", False),
    ("sample_submission.csv", "Soumission", "TableSoumission", True),
)



def inspect_csv(path: Path) -> tuple[int, int]:
    row_count = 0
    column_count = 0
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        for row in csv.reader(source):
            row_count += 1
            column_count = max(column_count, len(row))
    if not row_count or not column_count:
        raise ValueError(f"Le fichier CSV est vide : {path}")
    return row_count, column_count


def append_csv_to_sheet(path: Path, sheet, has_header: bool) -> tuple[int, int]:
    row_count, column_count = inspect_csv(path)
    if row_count + (not has_header) > EXCEL_MAX_ROWS:
        raise ValueError(f"Le fichier dépasse la limite de lignes Excel : {path}")
    if column_count > EXCEL_MAX_COLUMNS:
        raise ValueError(f"Le fichier dépasse la limite de colonnes Excel : {path}")

    with path.open("r", encoding="utf-8-sig", newline="") as source:
        rows = csv.reader(source)
        if has_header:
            headers = next(rows)
            if len(headers) != column_count or len(set(headers)) != len(headers):
                raise ValueError(f"Les en-têtes Excel sont vides ou dupliqués : {path}")
        else:
            headers = [f"champ_{index}" for index in range(1, column_count + 1)]

        header_cells = []
        for header in headers:
            cell = WriteOnlyCell(sheet, value=header)
            cell.font = Font(bold=True)
            header_cells.append(cell)
        sheet.append(header_cells)
        for row in rows:
            row = [value if value.strip() else MISSING_VALUE for value in row]
            row.extend([MISSING_VALUE] * (column_count - len(row)))
            sheet.append(row)

    table_ref = (
        f"A1:{get_column_letter(column_count)}"
        f"{row_count + (not has_header)}"
    )
    table = Table(
        displayName=next(item[2] for item in SOURCES if item[0] == path.name),
        ref=table_ref,
    )
    table.tableColumns = [
        TableColumn(id=index, name=header)
        for index, header in enumerate(headers, start=1)
    ]
    table.autoFilter = AutoFilter(ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="In write-only mode you must add table columns manually",
            category=UserWarning,
        )
        sheet.add_table(table)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    return row_count, column_count


def main() -> None:
    workbook = Workbook(write_only=True)
    for filename, sheet_name, _, has_header in SOURCES:
        source_path = DATA_DIR / filename
        if not source_path.is_file():
            raise FileNotFoundError(f"Fichier CSV introuvable : {source_path}")

        sheet = workbook.create_sheet(title=sheet_name)
        row_count, column_count = append_csv_to_sheet(
            source_path,
            sheet,
            has_header,
        )
        
        print(
            f"{filename} : {row_count} lignes, {column_count} colonnes "
            f"ajoutées à la feuille « {sheet_name} »."
        )

    workbook.save(OUTPUT_FILE)
    print(f"Classeur Excel créé : {OUTPUT_FILE}")



if __name__ == "__main__":
    main()
