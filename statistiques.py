import csv
import re
from collections import Counter, defaultdict
from pathlib import Path
from statistics import fmean, median

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_DIR = Path(__file__).resolve().parent
DATA_DIR = PROJECT_DIR / "data"
OUTPUT_DIR = DATA_DIR / "statistiques"
EXCEL_OUTPUT = OUTPUT_DIR / "statistiques_descriptives.xlsx"
TIME_MARKER = re.compile(r"t(\d+)", re.IGNORECASE)
DETAILS = re.compile(r"\([^)]*\)|<[^>]*>|\$[^$]*\$")


def action_type(value: str) -> str:
    action = DETAILS.sub("", value).strip()
    if action.endswith("1"):
        action = action[:-1].rstrip()
    return action


def read_dataset(filename: str, training: bool) -> dict:
    path = DATA_DIR / filename
    if not path.is_file():
        raise FileNotFoundError(f"Fichier de données introuvable : {path}")

    result = {
        "sessions": 0,
        "browser_sessions": Counter(),
        "browser_users": defaultdict(set),
        "session_action_counts": defaultdict(list),
        "session_durations": defaultdict(list),
        "session_time_rows": [],
        "user_browser_time": defaultdict(
            lambda: {"sessions": 0, "duration_seconds": 0}
        ),
        "action_occurrences": Counter(),
        "action_sessions": Counter(),
        "action_users": defaultdict(set),
        "user_sessions": Counter(),
        "user_browsers": defaultdict(Counter),
        "class_sessions": Counter(),
        "total_actions": 0,
    }

    with path.open("r", encoding="utf-8-sig", newline="") as source:
        for row_number, row in enumerate(csv.reader(source), start=1):
            if not row or not any(value.strip() for value in row):
                continue
            required_fields = 2 if training else 1
            if len(row) < required_fields:
                raise ValueError(
                    f"Ligne {row_number} invalide dans {path} : "
                    f"{required_fields} champs minimum attendus."
                )

            user = row[0].strip() if training else None
            browser = row[1].strip() if training else row[0].strip()
            events = row[2:] if training else row[1:]
            actions = Counter()
            duration = 0

            for value in events:
                value = value.strip()
                if not value:
                    continue
                time_match = TIME_MARKER.fullmatch(value)
                if time_match:
                    duration = max(duration, int(time_match.group(1)))
                    continue
                normalized_action = action_type(value)
                if not normalized_action:
                    continue
                actions[normalized_action] += 1

            result["sessions"] += 1
            session_number = result["sessions"]
            result["browser_sessions"][browser] += 1
            result["session_action_counts"][browser].append(sum(actions.values()))
            result["session_durations"][browser].append(duration)
            result["action_occurrences"].update(actions)
            result["action_sessions"].update(actions.keys())
            result["total_actions"] += sum(actions.values())

            if training:
                result["class_sessions"][user] += 1
                result["user_sessions"][user] += 1
                result["user_browsers"][user][browser] += 1
                result["browser_users"][browser].add(user)
                result["user_browser_time"][(user, browser)]["sessions"] += 1
                result["user_browser_time"][(user, browser)][
                    "duration_seconds"
                ] += duration
                result["session_time_rows"].append(
                    {
                        "session": session_number,
                        "utilisateur": user,
                        "navigateur": browser,
                        "derniere_borne_t_secondes": duration,
                        "duree_estimee_minutes": round(duration / 60, 2),
                    }
                )
                for action in actions:
                    result["action_users"][action].add(user)
            else:
                result["session_time_rows"].append(
                    {
                        "session": session_number,
                        "utilisateur": "",
                        "navigateur": browser,
                        "derniere_borne_t_secondes": duration,
                        "duree_estimee_minutes": round(duration / 60, 2),
                    }
                )

    return result


def write_csv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def add_excel_table(
    sheet,
    rows: list[dict],
    table_name: str,
    start_row: int = 1,
) -> int:
    if not rows:
        return start_row

    headers = list(rows[0])
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(start_row, column, header)
        cell.font = Font(bold=True)
    for row_number, row in enumerate(rows, start=start_row + 1):
        for column, header in enumerate(headers, start=1):
            sheet.cell(row_number, column, row[header])

    end_row = start_row + len(rows)
    table = Table(
        displayName=table_name,
        ref=f"A{start_row}:{sheet.cell(end_row, len(headers)).coordinate}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    return end_row + 1


def format_excel_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column in sheet.columns:
        column_letter = column[0].column_letter
        max_length = max(
            (len(str(cell.value)) for cell in column if cell.value is not None),
            default=10,
        )
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 55)


def write_excel(
    summary_rows: list[dict],
    train_browser_rows: list[dict],
    test_browser_rows: list[dict],
    user_rows: list[dict],
    action_rows: list[dict],
    session_time_rows: list[dict],
    user_browser_time_rows: list[dict],
) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resume"
    summary_sheet["A1"] = "STATISTIQUES DESCRIPTIVES - traces Copilote"
    summary_sheet["A1"].font = Font(bold=True, size=14)
    add_excel_table(summary_sheet, summary_rows, "TableResume", start_row=3)
    browser_start = add_excel_table(
        summary_sheet,
        train_browser_rows,
        "TableNavigateursResume",
        start_row=14,
    )
    add_excel_table(
        summary_sheet,
        action_rows[:10],
        "TableActionsTop10",
        start_row=browser_start + 1,
    )
    summary_sheet.freeze_panes = "A4"
    summary_sheet.sheet_view.showGridLines = False

    sheets = (
        ("Navigateurs train", train_browser_rows, "TableNavigateursTrain"),
        ("Navigateurs test", test_browser_rows, "TableNavigateursTest"),
        ("Temps par session", session_time_rows, "TableTempsSession"),
        (
            "Temps utilisateur",
            user_browser_time_rows,
            "TableTempsUtilisateur",
        ),
        ("Utilisateurs", user_rows, "TableUtilisateurs"),
        ("Actions", action_rows, "TableActions"),
    )
    for title, rows, table_name in sheets:
        sheet = workbook.create_sheet(title=title)
        add_excel_table(sheet, rows, table_name)
        format_excel_sheet(sheet)

    format_excel_sheet(summary_sheet)
    workbook.save(EXCEL_OUTPUT)


def describe_browser_stats(stats: dict, training: bool) -> list[dict]:
    result = []
    for browser, sessions in stats["browser_sessions"].most_common():
        durations = stats["session_durations"][browser]
        result.append(
            {
                "navigateur": browser,
                "sessions": sessions,
                "utilisateurs_distincts": (
                    len(stats["browser_users"][browser]) if training else ""
                ),
                "temps_total_heures": round(sum(durations) / 3600, 2),
                "temps_moyen_minutes_par_session": round(
                    fmean(durations) / 60, 2
                ),
            }
        )
    return result


def main() -> None:
    train = read_dataset("train.csv", training=True)
    test = read_dataset("test.csv", training=False)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    write_csv(
        OUTPUT_DIR / "navigateurs_train.csv",
        [
            "navigateur",
            "sessions",
            "utilisateurs_distincts",
            "temps_total_heures",
            "temps_moyen_minutes_par_session",
        ],
        describe_browser_stats(train, training=True),
    )
    write_csv(
        OUTPUT_DIR / "navigateurs_test.csv",
        [
            "navigateur",
            "sessions",
            "utilisateurs_distincts",
            "temps_total_heures",
            "temps_moyen_minutes_par_session",
        ],
        describe_browser_stats(test, training=False),
    )

    user_rows = []
    for user, sessions in train["user_sessions"].most_common():
        browsers = train["user_browsers"][user]
        user_rows.append(
            {
                "utilisateur": user,
                "sessions": sessions,
                "part_sessions_train_pct": round(
                    100 * sessions / train["sessions"], 2
                ),
                "navigateurs_utilises": " | ".join(sorted(browsers)),
                "n_navigateurs": len(browsers),
            }
        )
    write_csv(
        OUTPUT_DIR / "utilisateurs.csv",
        [
            "utilisateur",
            "sessions",
            "part_sessions_train_pct",
            "navigateurs_utilises",
            "n_navigateurs",
        ],
        user_rows,
    )

    user_browser_time_rows = []
    for (user, browser), values in sorted(train["user_browser_time"].items()):
        user_browser_time_rows.append(
            {
                "utilisateur": user,
                "navigateur": browser,
                "sessions": values["sessions"],
                "temps_total_minutes": round(
                    values["duration_seconds"] / 60, 2
                ),
                "temps_moyen_minutes_par_session": round(
                    values["duration_seconds"] / values["sessions"] / 60, 2
                ),
            }
        )
    write_csv(
        OUTPUT_DIR / "temps_utilisateurs_navigateurs.csv",
        [
            "utilisateur",
            "navigateur",
            "sessions",
            "temps_total_minutes",
            "temps_moyen_minutes_par_session",
        ],
        user_browser_time_rows,
    )
    write_csv(
        OUTPUT_DIR / "temps_sessions_train.csv",
        [
            "session",
            "utilisateur",
            "navigateur",
            "derniere_borne_t_secondes",
            "duree_estimee_minutes",
        ],
        train["session_time_rows"],
    )
    write_csv(
        OUTPUT_DIR / "temps_sessions_test.csv",
        [
            "session",
            "utilisateur",
            "navigateur",
            "derniere_borne_t_secondes",
            "duree_estimee_minutes",
        ],
        test["session_time_rows"],
    )

    action_rows = []
    for action, occurrences in train["action_occurrences"].most_common():
        action_rows.append(
            {
                "type_action": action,
                "occurrences": occurrences,
                "part_actions_pct": round(
                    100 * occurrences / train["total_actions"], 3
                ),
                "utilisateurs_distincts": len(train["action_users"][action]),
                "part_utilisateurs_train_pct": round(
                    100
                    * len(train["action_users"][action])
                    / len(train["user_sessions"]),
                    2,
                ),
                "sessions_concernees": train["action_sessions"][action],
                "part_sessions_train_pct": round(
                    100 * train["action_sessions"][action] / train["sessions"],
                    2,
                ),
            }
        )
    write_csv(
        OUTPUT_DIR / "actions.csv",
        [
            "type_action",
            "occurrences",
            "part_actions_pct",
            "utilisateurs_distincts",
            "part_utilisateurs_train_pct",
            "sessions_concernees",
            "part_sessions_train_pct",
        ],
        action_rows,
    )

    browser_rows = describe_browser_stats(train, training=True)
    top_browser = max(browser_rows, key=lambda row: row["sessions"])
    top_user_browser = max(
        browser_rows, key=lambda row: row["utilisateurs_distincts"]
    )
    all_durations = [
        duration
        for durations in train["session_durations"].values()
        for duration in durations
    ]
    total_connection_hours = sum(all_durations) / 3600
    summary = [
        "STATISTIQUES DESCRIPTIVES - traces Copilote",
        "",
        f"Sessions d'entrainement : {train['sessions']}",
        f"Sessions de test : {test['sessions']}",
        f"Utilisateurs distincts dans train : {len(train['user_sessions'])}",
        f"Navigateurs observes dans train : {len(train['browser_sessions'])}",
        f"Types d'action observes dans train : {len(train['action_occurrences'])}",
        f"Temps total estime des sessions train : {total_connection_hours:.2f} heures",
        "",
        (
            f"Navigateur avec le plus de sessions : {top_browser['navigateur']} "
            f"({top_browser['sessions']} sessions, "
            f"{top_browser['temps_total_heures']} heures cumulees)"
        ),
        (
            f"Navigateur utilise par le plus d'utilisateurs distincts : "
            f"{top_user_browser['navigateur']} "
            f"({top_user_browser['utilisateurs_distincts']} utilisateurs)"
        ),
        (
            f"Duree estimee par session : mediane "
            f"{median(all_durations) / 60:.2f} minutes, maximum "
            f"{max(all_durations) / 60:.2f} minutes"
        ),
        "",
        "Temps de connexion par navigateur (train) :",
    ]
    for row in browser_rows:
        summary.append(
            f"- {row['navigateur']} : {row['sessions']} sessions, "
            f"{row['utilisateurs_distincts']} utilisateurs, "
            f"{row['temps_total_heures']} heures cumulees, "
            f"{row['temps_moyen_minutes_par_session']} minutes/session"
        )
    summary.extend(["", "10 types d'action les plus frequents :"])
    for row in action_rows[:10]:
        summary.append(
            f"- {row['type_action']} : {row['occurrences']} occurrences "
            f"dans {row['sessions_concernees']} sessions, "
            f"utilisee par {row['utilisateurs_distincts']} utilisateurs"
        )

    summary_path = OUTPUT_DIR / "resume.txt"
    summary_path.write_text("\n".join(summary) + "\n", encoding="utf-8")
    summary_rows = [
        {"indicateur": "Sessions d'entrainement", "valeur": train["sessions"]},
        {"indicateur": "Sessions de test", "valeur": test["sessions"]},
        {
            "indicateur": "Utilisateurs distincts dans train",
            "valeur": len(train["user_sessions"]),
        },
        {
            "indicateur": "Navigateurs observes dans train",
            "valeur": len(train["browser_sessions"]),
        },
        {
            "indicateur": "Types d'action observes dans train",
            "valeur": len(train["action_occurrences"]),
        },
        {
            "indicateur": "Temps total estime train (heures)",
            "valeur": round(total_connection_hours, 2),
        },
        {
            "indicateur": "Duree mediane estimee par session (minutes)",
            "valeur": round(median(all_durations) / 60, 2),
        },
        {
            "indicateur": "Duree maximale estimee par session (minutes)",
            "valeur": round(max(all_durations) / 60, 2),
        },
    ]
    write_excel(
        summary_rows,
        browser_rows,
        describe_browser_stats(test, training=False),
        user_rows,
        action_rows,
        train["session_time_rows"],
        user_browser_time_rows,
    )
    print("\n".join(summary))
    print(f"\nTableaux CSV et Excel enregistrés dans : {OUTPUT_DIR}")
    print(f"Classeur Excel : {EXCEL_OUTPUT}")


if __name__ == "__main__":
    main()
